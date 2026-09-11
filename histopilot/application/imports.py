"""Bounded real CSV/XLSX imports into immutable, folder-owned dataset records.

Discovery inspects file identity and size, never WSI pixels. Missing patient IDs
remain unresolved unless slide-level grouping is explicitly selected and recorded.
"""

import base64
import binascii
import csv
import hashlib
import io
import json
import math
import os
import re
import stat
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError

from pydantic import ValidationError

from histopilot.schemas.imports import AttributeMapping, ImportSpec, InspectRequest, TableSource
from histopilot.storage.filesystem import FilesystemError, LocalFilesystem
from histopilot.storage.project_lock import StorageError
from histopilot.storage.scientific import MAX_ARTIFACT_BYTES, ScientificStore

MAX_TABLE_BYTES = 16 * 1024 * 1024
MAX_UPLOAD_BYTES = 256 * 1024
MAX_ROWS = 20000
MAX_COLUMNS = 128
MAX_FILES = 10000
MAX_ENTRIES = 20000
PARSER_VERSION = "histopilot-tables-v1"
SLIDE_EXTENSIONS = {
    ".svs",
    ".tif",
    ".tiff",
    ".ndpi",
    ".mrxs",
    ".scn",
    ".vms",
    ".vmu",
    ".bif",
    ".sdpc",
}
RESERVED = {"Slide_ID", "Patient_ID", "Slide_Path", "slideId", "patientId", "slidePath"}


def _error(message, code="IMPORT_INVALID", status=422):
    return StorageError(message, code, status)


def _json(value):
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    ).encode("utf-8")


class _Findings:
    def __init__(self):
        self.entries = {}

    def add(self, code, message, *, severity="error", example=None, count=1):
        key = (severity, code, message)
        value = self.entries.setdefault(
            key,
            {"severity": severity, "code": code, "message": message, "count": 0, "examples": []},
        )
        value["count"] += count
        if example is not None and len(value["examples"]) < 5:
            value["examples"].append(str(example))

    def merge(self, findings):
        for item in findings:
            self.add(
                item["code"], item["message"], severity=item["severity"], count=item.get("count", 1)
            )
            self.entries[(item["severity"], item["code"], item["message"])]["examples"] = item.get(
                "examples", []
            )[:5]

    def result(self):
        return [self.entries[key] for key in sorted(self.entries)]


class ImportService:
    def __init__(self, store: ScientificStore, filesystem: LocalFilesystem):
        self.store = store
        self.filesystem = filesystem

    def _source_bytes(self, source: TableSource):
        filename = source.filename
        path = None
        if source.path is not None:
            candidate = Path(source.path)
            if not candidate.is_absolute() or "\x00" in source.path:
                raise FilesystemError("Choose an absolute server metadata file path.")
            try:
                path = candidate.resolve(strict=True)
            except (OSError, RuntimeError) as error:
                raise FilesystemError("The metadata file cannot be resolved.", 404) from error
            if not self.filesystem._contains(path):
                raise FilesystemError("The metadata file is outside configured data roots.", 403)
            descriptor = None
            try:
                descriptor = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
                before = os.fstat(descriptor)
                if not stat.S_ISREG(before.st_mode):
                    raise _error(
                        "Select a regular CSV or XLSX metadata file.", "TABLE_FILE_REQUIRED"
                    )
                if before.st_size > MAX_TABLE_BYTES:
                    raise _error("Metadata files must be at most 16 MiB.", "TABLE_TOO_LARGE", 413)
                with os.fdopen(descriptor, "rb", closefd=False) as handle:
                    content = handle.read(MAX_TABLE_BYTES + 1)
                after = os.fstat(descriptor)
                if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
                    raise _error(
                        "The metadata file changed while being read.", "SOURCE_CHANGED", 409
                    )
                if len(content) > MAX_TABLE_BYTES:
                    raise _error("Metadata files must be at most 16 MiB.", "TABLE_TOO_LARGE", 413)
            except OSError as error:
                raise FilesystemError("The metadata file cannot be read.", 403) from error
            finally:
                if descriptor is not None:
                    os.close(descriptor)
            filename = path.name
        else:
            try:
                content = base64.b64decode(source.contentBase64, validate=True)
            except (ValueError, binascii.Error) as error:
                raise _error("The uploaded file is not valid base64.", "UPLOAD_INVALID") from error
            if len(content) > MAX_UPLOAD_BYTES:
                raise _error(
                    "Inline metadata uploads must be at most 256 KiB; use a server file path for larger tables.",
                    "UPLOAD_TOO_LARGE",
                    413,
                )
        extension = Path(filename).suffix.lower()
        if extension not in {".csv", ".xlsx"}:
            raise _error("Select a CSV or XLSX metadata file.", "TABLE_FORMAT_UNSUPPORTED")
        return content, {
            "filename": filename,
            "path": str(path) if path else None,
            "format": extension[1:],
            "sha256": hashlib.sha256(content).hexdigest(),
            "sizeBytes": len(content),
        }

    def _table(self, source: TableSource):
        content, provenance = self._source_bytes(source)
        findings = _Findings()
        numeric_columns, formula_columns = set(), set()
        error_cells = defaultdict(set)
        sheets, sheet = [], None
        workbook = None
        try:
            if provenance["format"] == "csv":
                if source.sheet is not None:
                    raise _error("CSV files do not have worksheets.", "SHEET_INVALID")
                try:
                    text = content.decode("utf-8-sig")
                except UnicodeError as error:
                    raise _error(
                        "CSV metadata must use UTF-8 encoding.", "TABLE_ENCODING"
                    ) from error
                rows = csv.reader(io.StringIO(text, newline=""), strict=True)
                raw_rows = ((list(row), [], []) for row in rows)
            else:
                from openpyxl import load_workbook
                from openpyxl.utils.exceptions import InvalidFileException

                try:
                    with zipfile.ZipFile(io.BytesIO(content)) as archive:
                        if (
                            len(archive.infolist()) > 2000
                            or sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024
                        ):
                            raise _error(
                                "The workbook expands beyond the supported metadata size.",
                                "TABLE_TOO_LARGE",
                                413,
                            )
                    workbook = load_workbook(
                        io.BytesIO(content), read_only=True, data_only=False, keep_links=False
                    )
                except StorageError:
                    raise
                except (
                    zipfile.BadZipFile,
                    InvalidFileException,
                    ValueError,
                    KeyError,
                    OSError,
                    ParseError,
                ) as error:
                    raise _error(
                        "The XLSX workbook could not be parsed.", "TABLE_INVALID"
                    ) from error
                sheets = workbook.sheetnames
                sheet = source.sheet or (sheets[0] if sheets else None)
                if sheet not in sheets:
                    raise _error("Select an existing worksheet.", "SHEET_INVALID")
                worksheet = workbook[sheet]
                if worksheet.max_row and worksheet.max_row > MAX_ROWS + 1:
                    raise _error("The worksheet exceeds 20,000 data rows.", "TABLE_ROW_LIMIT", 413)
                if worksheet.max_column and worksheet.max_column > MAX_COLUMNS:
                    raise _error("The worksheet exceeds 128 columns.", "TABLE_COLUMN_LIMIT", 413)

                def cells():
                    for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                        values, numeric, formulas = [], [], []
                        for index, cell in enumerate(row):
                            value = cell.value
                            if cell.data_type == "f":
                                formulas.append(index)
                            elif cell.data_type == "e" and row_number > 1:
                                error_cells[row_number].add(index)
                            if isinstance(value, (int, float)) and not isinstance(value, bool):
                                numeric.append(index)
                            if isinstance(value, (datetime, date)):
                                value = value.isoformat()
                            elif value is not None:
                                value = str(value)
                            values.append(value)
                        yield values, numeric, formulas

                raw_rows = cells()
            first = next(raw_rows, None)
            if first is None:
                raise _error("The table has no header row.", "TABLE_EMPTY")
            headers = first[0]
            if (
                not headers
                or len(headers) > MAX_COLUMNS
                or any(not isinstance(header, str) or not header for header in headers)
            ):
                raise _error(
                    "Use one nonempty header per column, up to 128 columns.",
                    "TABLE_HEADERS_INVALID",
                )
            if len(set(headers)) != len(headers):
                raise _error(
                    "Duplicate column headers must be resolved before mapping.",
                    "TABLE_HEADERS_DUPLICATE",
                )
            if any(len(header) > 256 for header in headers):
                raise _error(
                    "Column headers must be at most 256 characters.", "TABLE_HEADERS_INVALID"
                )
            records = []
            for row_number, (values, numeric, formulas) in enumerate(raw_rows, start=2):
                if row_number > MAX_ROWS + 1:
                    raise _error(
                        "Metadata tables may contain at most 20,000 rows.", "TABLE_ROW_LIMIT", 413
                    )
                if not values or all(value in {None, ""} for value in values):
                    continue
                if len(values) != len(headers):
                    raise _error(
                        f"Row {row_number} does not match the header width.", "TABLE_ROW_WIDTH"
                    )
                if any(value is not None and len(value) > 65536 for value in values):
                    raise _error(
                        "A metadata cell exceeds the supported size.", "TABLE_CELL_LIMIT", 413
                    )
                numeric_columns.update(headers[index] for index in numeric)
                formula_columns.update(headers[index] for index in formulas)
                records.append(
                    {
                        "row": row_number,
                        "values": dict(zip(headers, values, strict=True)),
                        "errors": [
                            headers[index] for index in sorted(error_cells.get(row_number, ()))
                        ],
                    }
                )
            if not records:
                raise _error("The table contains no data rows.", "TABLE_EMPTY")
            if formula_columns:
                findings.add(
                    "FORMULA_CELLS",
                    "Formula cells are preserved as text and never evaluated; mapped formula fields must be replaced with reviewed values.",
                    severity="warning",
                    count=len(formula_columns),
                )
            if any(error_cells.values()):
                findings.add(
                    "SPREADSHEET_ERROR_CELLS",
                    "Spreadsheet error cells are preserved as source evidence. Replace mapped errors with reviewed values or explicitly map their tokens to missing.",
                    severity="warning",
                    count=sum(len(columns) for columns in error_cells.values()),
                )
            provenance["sheet"] = sheet
            return {
                "headers": headers,
                "sheets": sheets,
                "sheet": sheet,
                "rows": records,
                "provenance": provenance,
                "bytes": content,
                "numeric": numeric_columns,
                "formulas": formula_columns,
                "findings": findings.result(),
            }
        except (csv.Error, ParseError) as error:
            raise _error("The metadata table could not be parsed.", "TABLE_INVALID") from error
        finally:
            if workbook is not None:
                workbook.close()

    def inspect(self, spec: InspectRequest | TableSource):
        table = self._table(spec.source if isinstance(spec, InspectRequest) else spec)
        column_summaries = {}
        for header in table["headers"]:
            # Summarize the entire parsed source; the displayed rows below are only
            # a sample. Missing-value mappings have not been applied at this stage.
            counts = Counter(
                row["values"][header]
                for row in table["rows"]
                if row["values"][header] not in {None, ""}
            )
            column_summaries[header] = {
                # Counter preserves source encounter order when frequencies tie.
                "examples": [value for value, _ in counts.most_common(4)],
                "distinctCount": len(counts),
                "missingCount": len(table["rows"]) - counts.total(),
            }
        return {
            "headers": table["headers"],
            "sheets": table["sheets"],
            "sheet": table["sheet"],
            "rows": [row["values"] for row in table["rows"][:20]],
            "rowCount": len(table["rows"]),
            "columnSummaries": column_summaries,
            "fingerprint": table["provenance"]["sha256"],
            "findings": table["findings"],
        }

    def _scan(self, spec: ImportSpec, findings: _Findings):
        if spec.slideRoot is None:
            if not spec.includeMissingSlides:
                findings.add(
                    "SLIDE_SOURCE_REQUIRED",
                    "Choose a slide folder or explicitly retain rows without slide files.",
                )
            return []
        root = self.filesystem.directory(spec.slideRoot)
        inventory, pending, encountered, file_count = [], [root], 0, 0
        while pending:
            current = pending.pop()
            try:
                with os.scandir(current) as listing:
                    for entry in listing:
                        encountered += 1
                        if encountered > MAX_ENTRIES:
                            raise _error(
                                "The source tree exceeds the bounded scan limit; choose a smaller root.",
                                "SCAN_LIMIT",
                                413,
                            )
                        path = Path(entry.path)
                        try:
                            resolved = path.resolve(strict=True)
                            if not self.filesystem._contains(resolved):
                                findings.add(
                                    "SOURCE_PATH_ESCAPE",
                                    "A discovered source escapes configured data roots.",
                                    example=path.relative_to(root),
                                )
                                continue
                            info = resolved.stat()
                            if stat.S_ISDIR(info.st_mode):
                                if spec.recursive:
                                    if path.is_symlink():
                                        findings.add(
                                            "SYMLINK_DIRECTORY",
                                            "Recursive import requires physical directories; a linked subdirectory needs an explicit source choice.",
                                            example=path.relative_to(root),
                                        )
                                    else:
                                        pending.append(path)
                                continue
                            file_count += 1
                            if file_count > MAX_FILES:
                                raise _error(
                                    "The source tree exceeds 10,000 files; choose a smaller root.",
                                    "SCAN_LIMIT",
                                    413,
                                )
                            if path.suffix.lower() not in SLIDE_EXTENSIONS:
                                continue
                            if not stat.S_ISREG(info.st_mode):
                                findings.add(
                                    "SLIDE_FILE_INVALID",
                                    "A slide entry is not a regular file.",
                                    example=path.relative_to(root),
                                )
                                continue
                            if info.st_size == 0:
                                findings.add(
                                    "SLIDE_EMPTY",
                                    "A discovered slide file is empty.",
                                    example=path.relative_to(root),
                                )
                            inventory.append(
                                {
                                    "slideId": path.stem,
                                    "path": str(resolved),
                                    "relativePath": path.relative_to(root).as_posix(),
                                    "sizeBytes": info.st_size,
                                    "mtimeNs": info.st_mtime_ns,
                                    "device": info.st_dev,
                                    "inode": info.st_ino,
                                }
                            )
                        except (OSError, RuntimeError):
                            findings.add(
                                "SLIDE_UNREADABLE",
                                "A source entry cannot be inspected; the inventory is incomplete.",
                                example=path.relative_to(root),
                            )
            except OSError as error:
                raise FilesystemError(
                    "A selected slide directory cannot be scanned.", 403
                ) from error
        return sorted(inventory, key=lambda item: (item["slideId"], item["relativePath"]))

    @staticmethod
    def _value(value, missing):
        return None if value is None or value in missing else value

    def _identifier(self, value, missing, findings, *, column, row):
        # Attribute missing-value policies may deliberately preserve empty strings.
        # An empty identity can never establish a patient group or a join key.
        value = self._value(value, missing)
        if value is None or value == "":
            return None
        if not value.strip():
            findings.add(
                "IDENTIFIER_BLANK",
                f"Identifiers in {column} contain only whitespace; replace them with reviewed IDs or explicitly mark that token as missing.",
                example=f"row {row}",
            )
            return None
        if value != value.strip():
            findings.add(
                "IDENTIFIER_WHITESPACE",
                "Identifiers contain leading or trailing whitespace; values are preserved without implicit normalization.",
                severity="warning",
                example=value,
            )
        return value

    def _mapped_errors(self, table, identities, mappings, missing, findings):
        policies = defaultdict(list)
        for column in identities:
            if column is not None:
                policies[column].append(missing)
        for field in mappings:
            policies[field.sourceColumn].append(
                field.missingValues if field.missingValues is not None else missing
            )
        for row in table["rows"]:
            for column in row["errors"]:
                if any(
                    self._value(row["values"][column], policy) is not None
                    for policy in policies[column]
                ):
                    findings.add(
                        "MAPPED_SPREADSHEET_ERROR",
                        "A mapped field contains a spreadsheet error. Replace it with a reviewed value or explicitly declare that token missing before freezing.",
                        example=f"{column}, row {row['row']}",
                    )

    def _attributes(self, mappings, values, missing, findings, row_id):
        result = {}
        for field in mappings:
            value = self._value(
                values.get(field.sourceColumn),
                field.missingValues if field.missingValues is not None else missing,
            )
            result[field.key] = value
            if value is None:
                continue
            valid = True
            try:
                if field.type == "integer":
                    valid = re.fullmatch(r"[+-]?[0-9]+", value) is not None
                elif field.type == "decimal":
                    valid = Decimal(value).is_finite() and math.isfinite(float(value))
                elif field.type == "boolean":
                    valid = value.lower() in {"true", "false", "0", "1"}
                elif field.type == "date":
                    try:
                        date.fromisoformat(value)
                    except ValueError:
                        datetime.fromisoformat(value)
                elif (
                    field.type in {"categorical", "ordered_categorical"}
                    and field.categories is not None
                ):
                    valid = value in field.categories
            except (ValueError, InvalidOperation):
                valid = False
            if not valid:
                findings.add(
                    "ATTRIBUTE_TYPE_INVALID",
                    f"Values in {field.sourceColumn} do not match its declared type or categories.",
                    example=row_id,
                )
        return result

    def _draft(self, draft_id, expected_revision):
        if type(expected_revision) is not int or expected_revision < 1:
            raise _error("Supply a positive draft revision.", "INVALID_REVISION")
        draft = self.store.get_draft(draft_id)
        if draft["revision"] != expected_revision:
            raise _error(
                "The import draft changed. Reload before previewing or freezing.",
                "REVISION_CONFLICT",
                409,
            )
        if draft["status"] != "editable":
            raise _error("This import draft is already frozen.", "DRAFT_FROZEN", 409)
        payload = draft["payload"]
        if (
            draft["kind"] != "import"
            or set(payload) != {"type", "spec"}
            or payload["type"] != "dataset-import"
        ):
            raise _error("Use a dataset-import draft for this operation.", "IMPORT_DRAFT_REQUIRED")
        try:
            return draft, ImportSpec.model_validate(payload["spec"])
        except ValidationError as error:
            errors = error.errors(include_input=False, include_url=False)
            details = [
                f"{'.'.join(str(part) for part in item['loc']) or 'import'}: "
                f"{item['msg'].removeprefix('Value error, ')}"
                for item in errors[:8]
            ]
            if len(errors) > 8:
                details.append(f"{len(errors) - 8} more fields need correction.")
            raise _error(
                "Review the import source and mappings. " + " ".join(details),
                "IMPORT_SPEC_INVALID",
            ) from error

    def _build(self, spec: ImportSpec):
        primary = self._table(spec.source)
        secondary = self._table(spec.patientSource) if spec.patientSource is not None else None
        findings = _Findings()
        findings.merge(primary["findings"])
        if secondary:
            findings.merge(secondary["findings"])
        main_fields = spec.attributes
        if main_fields is None:
            if any(
                len(column) > 128
                for column in primary["headers"]
                if column not in {spec.slideIdColumn, spec.patientIdColumn}
            ):
                raise _error(
                    "Map long source headers to explicit attribute keys of at most 128 characters.",
                    "ATTRIBUTE_KEY_REQUIRED",
                )
            main_fields = [
                AttributeMapping(key=column, sourceColumn=column)
                for column in primary["headers"]
                if column not in {spec.slideIdColumn, spec.patientIdColumn}
            ]
        all_fields = [*main_fields, *spec.patientAttributes]
        dictionary = [field.model_dump(exclude_none=True) for field in all_fields]
        keys = [field.key for field in all_fields]
        if len(set(keys)) != len(keys) or any(key in RESERVED for key in keys):
            findings.add(
                "ATTRIBUTE_KEY_CONFLICT",
                "Attribute keys must be unique and cannot use reserved identity fields.",
            )
        required_main = {spec.slideIdColumn, *(field.sourceColumn for field in main_fields)}
        if spec.patientIdColumn:
            required_main.add(spec.patientIdColumn)
        if required_main - set(primary["headers"]):
            raise _error(
                "One or more selected columns are absent from the main table.",
                "MAPPING_COLUMN_MISSING",
            )
        if spec.patientIdColumn == spec.slideIdColumn:
            findings.add(
                "PATIENT_MAPPING_REQUIRES_CROSSWALK",
                "A slide identifier cannot also be asserted as patient identity; supply a separately reviewed patient field or crosswalk.",
            )
        mapped_primary = required_main
        if mapped_primary & primary["formulas"]:
            findings.add(
                "MAPPED_FORMULA", "A mapped primary-table field contains unevaluated formulas."
            )
        self._mapped_errors(
            primary,
            {spec.slideIdColumn, spec.patientIdColumn},
            main_fields,
            spec.missingValues,
            findings,
        )
        if {spec.slideIdColumn, spec.patientIdColumn} & primary["numeric"]:
            findings.add(
                "NUMERIC_IDENTIFIER",
                "Numeric spreadsheet identifiers retain the stored value; formatting or previously lost zeros cannot establish the original ID.",
                severity="warning",
            )
        inventory = self._scan(spec, findings)
        matches = defaultdict(list)
        physical = defaultdict(list)
        for item in inventory:
            matches[item["slideId"]].append(item)
            physical[(item["device"], item["inode"])].append(item)
        for slide_id, items in matches.items():
            if len(items) > 1:
                findings.add(
                    "SLIDE_MATCH_AMBIGUOUS",
                    "A filename stem resolves to multiple slide files.",
                    example=slide_id,
                )
        for items in physical.values():
            if len(items) > 1:
                findings.add(
                    "DUPLICATE_SLIDE_ALIAS",
                    "Multiple inventory paths identify the same physical file.",
                    example=items[0]["relativePath"],
                )
        patient_rows, secondary_used = {}, set()
        if secondary:
            if spec.patientSourceKind == "patients" and not spec.patientIdColumn:
                findings.add(
                    "PATIENT_JOIN_REQUIRED",
                    "A patient-keyed table requires an explicit patient ID column in the main slide table.",
                )
            if (
                spec.patientSourceKind == "crosswalk"
                and spec.patientSourceSlideIdColumn == spec.patientSourcePatientIdColumn
            ):
                findings.add(
                    "PATIENT_MAPPING_REQUIRES_CROSSWALK",
                    "A crosswalk needs distinct slide and patient source columns.",
                )
            required = {
                spec.patientSourcePatientIdColumn,
                *(field.sourceColumn for field in spec.patientAttributes),
            }
            key_column = (
                spec.patientSourceSlideIdColumn
                if spec.patientSourceKind == "crosswalk"
                else spec.patientSourcePatientIdColumn
            )
            required.add(key_column)
            if required - set(secondary["headers"]):
                raise _error(
                    "One or more selected columns are absent from the patient source.",
                    "MAPPING_COLUMN_MISSING",
                )
            if required & secondary["formulas"]:
                findings.add(
                    "MAPPED_FORMULA", "A mapped patient-source field contains unevaluated formulas."
                )
            self._mapped_errors(
                secondary,
                {key_column, spec.patientSourcePatientIdColumn},
                spec.patientAttributes,
                spec.missingValues,
                findings,
            )
            if {key_column, spec.patientSourcePatientIdColumn} & secondary["numeric"]:
                findings.add(
                    "NUMERIC_IDENTIFIER",
                    "Numeric spreadsheet identifiers retain the stored value; formatting or previously lost zeros cannot establish the original ID.",
                    severity="warning",
                )
            for row in secondary["rows"]:
                key = self._identifier(
                    row["values"][key_column],
                    spec.missingValues,
                    findings,
                    column=key_column,
                    row=row["row"],
                )
                patient = (
                    key
                    if key_column == spec.patientSourcePatientIdColumn
                    else self._identifier(
                        row["values"][spec.patientSourcePatientIdColumn],
                        spec.missingValues,
                        findings,
                        column=spec.patientSourcePatientIdColumn,
                        row=row["row"],
                    )
                )
                if key is None or patient is None:
                    findings.add(
                        "PATIENT_SOURCE_KEY_MISSING",
                        "Patient-source rows need an explicit join key and patient identifier.",
                        example=f"row {row['row']}",
                    )
                elif key in patient_rows:
                    findings.add(
                        "PATIENT_SOURCE_KEY_DUPLICATE",
                        "The patient source contains duplicate join keys.",
                        example=key,
                    )
                else:
                    patient_rows[key] = row["values"]
        records, exclusions, source_ids = [], [], set()
        record_bytes = 2
        for row in primary["rows"]:
            values = row["values"]
            slide_id = self._identifier(
                values[spec.slideIdColumn],
                spec.missingValues,
                findings,
                column=spec.slideIdColumn,
                row=row["row"],
            )
            if slide_id is None:
                findings.add(
                    "SLIDE_ID_MISSING",
                    "Every included source row requires a Slide_ID.",
                    example=f"row {row['row']}",
                )
                continue
            if slide_id in source_ids:
                findings.add(
                    "SLIDE_ID_DUPLICATE",
                    "The main table contains duplicate Slide_ID values.",
                    example=slide_id,
                )
                continue
            source_ids.add(slide_id)
            patient_id = (
                self._identifier(
                    values[spec.patientIdColumn],
                    spec.missingValues,
                    findings,
                    column=spec.patientIdColumn,
                    row=row["row"],
                )
                if spec.patientIdColumn
                else None
            )
            patient_id_source = "source" if patient_id is not None else "unresolved"
            join_key = slide_id if spec.patientSourceKind == "crosswalk" else patient_id
            patient_values = patient_rows.get(join_key)
            if patient_values:
                resolved_patient = self._value(
                    patient_values[spec.patientSourcePatientIdColumn], spec.missingValues
                )
                if patient_id is not None and patient_id != resolved_patient:
                    findings.add(
                        "PATIENT_LINK_CONFLICT",
                        "The main table and patient crosswalk disagree on patient identity.",
                        example=slide_id,
                    )
                else:
                    patient_id = resolved_patient
                    if spec.patientSourceKind == "crosswalk":
                        patient_id_source = "crosswalk"
            files = matches.get(slide_id, [])
            if not files and not spec.includeMissingSlides:
                exclusions.append(
                    {"slideId": slide_id, "sourceRow": row["row"], "reason": "SLIDE_FILE_MISSING"}
                )
                continue
            if patient_values:
                secondary_used.add(join_key)
            elif secondary and join_key is not None:
                findings.add(
                    "PATIENT_SOURCE_MATCH_MISSING",
                    "Some included slides have no matching row in the selected patient source; review their patient links and missing patient-source attributes.",
                    severity="warning",
                    example=slide_id,
                )
            # Resolve true patient identities and joins before applying the opted-in
            # fallback. A slide identifier must never join a patient attribute table.
            if patient_id is None and spec.patientIdFallback == "slide_id":
                patient_id = slide_id
                patient_id_source = "slide_fallback"
            attributes = self._attributes(
                main_fields, values, spec.missingValues, findings, slide_id
            )
            attributes.update(
                self._attributes(
                    spec.patientAttributes,
                    patient_values or {},
                    spec.missingValues,
                    findings,
                    slide_id,
                )
            )
            record = {
                "slideId": slide_id,
                "patientId": patient_id,
                "patientIdSource": patient_id_source,
                "slidePath": files[0]["path"] if len(files) == 1 else None,
                "attributes": attributes,
            }
            record_bytes += len(_json(record)) + bool(records)
            if record_bytes > MAX_ARTIFACT_BYTES:
                raise _error(
                    "Joined records exceed the supported artifact size; reduce the selected scalar fields or source population.",
                    "IMPORT_TOO_LARGE",
                    413,
                )
            records.append(record)
        patient_spellings = defaultdict(set)
        for record in records:
            if record["patientId"] is not None:
                patient_spellings[record["patientId"].strip()].add(record["patientId"])
        for spellings in patient_spellings.values():
            if len(spellings) > 1:
                findings.add(
                    "PATIENT_ID_WHITESPACE_COLLISION",
                    "Patient IDs differ only in leading or trailing whitespace and would form separate split groups. Reconcile these identities explicitly before freezing.",
                    example=", ".join(repr(value) for value in sorted(spellings)),
                )
        verified_patients = {
            record["patientId"]
            for record in records
            if record["patientId"] is not None and record["patientIdSource"] != "slide_fallback"
        }
        fallback_ids = {
            record["patientId"]
            for record in records
            if record["patientIdSource"] == "slide_fallback"
        }
        for collision in sorted(fallback_ids & verified_patients):
            findings.add(
                "PATIENT_ID_FALLBACK_COLLISION",
                "A fallback Slide_ID matches a supplied Patient_ID. Add a verified patient mapping for this slide or leave its Patient_ID unresolved before freezing; these identities cannot safely be grouped together.",
                example=collision,
            )
        patient_values = defaultdict(set)
        for record in records:
            if record["patientId"] is None:
                continue
            for field in all_fields:
                if field.owner == "patient":
                    patient_values[(record["patientId"], field.key)].add(
                        record["attributes"].get(field.key)
                    )
        for (patient, key), values in patient_values.items():
            known = values - {None}
            if len(known) > 1:
                findings.add(
                    "PATIENT_ATTRIBUTE_CONFLICT",
                    f"Patient-owned attribute {key} has conflicting nonmissing values.",
                    example=patient,
                )
            elif known and None in values:
                findings.add(
                    "PATIENT_ATTRIBUTE_INCOMPLETE",
                    f"Patient-owned attribute {key} mixes known and missing values; explicitly reconcile the repeated patient value before freezing.",
                    example=patient,
                )
        if not records:
            findings.add("DATASET_EMPTY", "The reviewed mapping includes no slide records.")
        unmatched = [item for item in inventory if item["slideId"] not in source_ids]
        if exclusions:
            findings.add(
                "ROWS_WITHOUT_SLIDES",
                "Rows without matched slide files are excluded by the saved inclusion rule.",
                severity="warning",
                count=len(exclusions),
            )
        if unmatched:
            findings.add(
                "FILES_WITHOUT_ROWS",
                "Discovered slide files without metadata rows are retained in the unmatched inventory.",
                severity="warning",
                count=len(unmatched),
            )
        if secondary and set(patient_rows) - secondary_used:
            findings.add(
                "PATIENT_SOURCE_UNMATCHED",
                "Patient-source rows are not linked to the selected main table.",
                severity="warning",
                count=len(set(patient_rows) - secondary_used),
            )
        unlinked = sum(record["patientId"] is None for record in records)
        if unlinked:
            findings.add(
                "PATIENT_ID_UNRESOLVED",
                "Patient identity remains unresolved; patient-grouped execution is unavailable.",
                severity="warning",
                count=unlinked,
            )
        fallback_count = sum(record["patientIdSource"] == "slide_fallback" for record in records)
        if fallback_count:
            findings.add(
                "PATIENT_ID_SLIDE_FALLBACK",
                "Patient_ID was unresolved for these slides and explicitly falls back to Slide_ID. Each fallback slide is a separate grouping unit; patient identity is not verified and slides from the same patient may enter different splits.",
                severity="warning",
                count=fallback_count,
            )
        retained_without_slides = sum(record["slidePath"] is None for record in records)
        if retained_without_slides:
            findings.add(
                "SLIDES_UNAVAILABLE",
                "Included records have no unique available slide path; feature-only use requires separate feature validation.",
                severity="warning",
                count=retained_without_slides,
            )
        records.sort(key=lambda record: record["slideId"])
        summary = {
            "sourceRowCount": len(primary["rows"]),
            "slideCount": len(records),
            "mappedPatientCount": len(
                {record["patientId"] for record in records if record["patientId"] is not None}
            ),
            "verifiedPatientCount": len(verified_patients),
            "fallbackSlideCount": fallback_count,
            "unlinkedSlideCount": unlinked,
            "matchedSlideCount": len(records) - retained_without_slides,
            "missingSlideCount": len(exclusions) + retained_without_slides,
            "unmatchedFileCount": len(unmatched),
            "excludedRowCount": len(exclusions),
            "scannedFileCount": len(inventory),
        }
        provenance = {
            "parserVersion": PARSER_VERSION,
            "sources": {"main": primary["provenance"]},
            "mapping": spec.model_dump(exclude_none=True),
            "inventoryVerification": "file-identity-size-mtime-only",
        }
        # Upload bytes are snapshotted separately; do not duplicate them in manifests.
        for source_key in ("source", "patientSource"):
            if source_key in provenance["mapping"]:
                provenance["mapping"][source_key].pop("contentBase64", None)
        source_artifacts = {f"sources/main.{primary['provenance']['format']}": primary["bytes"]}
        if secondary:
            provenance["sources"]["patient"] = secondary["provenance"]
            source_artifacts[f"sources/patient.{secondary['provenance']['format']}"] = secondary[
                "bytes"
            ]
        result = {
            "summary": summary,
            "dictionary": dictionary,
            "records": records,
            "findings": findings.result(),
            "provenance": provenance,
            "inventory": inventory,
            "exclusions": exclusions,
        }
        preview_hash = hashlib.sha256(_json(result)).hexdigest()
        artifacts = {
            "records.json": _json(records),
            "dictionary.json": _json(dictionary),
            "inventory.json": _json(inventory),
            "exclusions.json": _json(exclusions),
            "provenance.json": _json(provenance),
            **source_artifacts,
        }
        if any(len(value) > MAX_ARTIFACT_BYTES for value in artifacts.values()):
            raise _error(
                "The materialized table exceeds the supported artifact size.",
                "IMPORT_TOO_LARGE",
                413,
            )
        manifest = {
            "kind": "dataset",
            "schemaVersion": 1,
            "summary": summary,
            "dictionary": dictionary,
            "previewHash": preview_hash,
            "provenance": provenance,
            "findings": result["findings"],
        }
        if spec.parentId is not None:
            manifest["parentId"] = spec.parentId
        return result, manifest, artifacts

    def preview(self, draft_id, expected_revision):
        draft, spec = self._draft(draft_id, expected_revision)
        result, manifest, _ = self._build(spec)
        return {
            "draftId": draft_id,
            "revision": draft["revision"],
            "previewHash": manifest["previewHash"],
            "canFreeze": not any(item["severity"] == "error" for item in result["findings"]),
            "findings": result["findings"],
            "summary": result["summary"],
            "dictionary": result["dictionary"],
            "records": result["records"][:200],
            "recordsTruncated": len(result["records"]) > 200,
        }

    def freeze(
        self, draft_id, expected_revision, preview_hash, operation_id, *, version_label=None
    ):
        if type(expected_revision) is not int or expected_revision < 1:
            raise _error("Supply a positive draft revision.", "INVALID_REVISION")
        previous = self.store.replay_dataset_publication(
            operation_id,
            draft_id=draft_id,
            expected_revision=expected_revision,
            preview_hash=preview_hash,
            version_label=version_label,
        )
        if previous is not None:
            return previous
        _, spec = self._draft(draft_id, expected_revision)
        result, manifest, artifacts = self._build(spec)
        if manifest["previewHash"] != preview_hash:
            raise _error(
                "Sources or import choices changed since preview. Review a new preview before freezing.",
                "PREVIEW_STALE",
                409,
            )
        if any(item["severity"] == "error" for item in result["findings"]):
            raise _error("Resolve the blocking import findings before freezing.", "IMPORT_BLOCKED")
        return self.store.publish_dataset(
            draft_id,
            expected_revision=expected_revision,
            manifest=manifest,
            artifacts=artifacts,
            operation_id=operation_id,
            version_label=version_label,
        )

    def records(self, dataset_id):
        dataset = self.store.get_dataset(dataset_id)
        if dataset["manifest"].get("kind") != "dataset":
            raise _error("This record is not an imported dataset.", "DATASET_KIND_INVALID")
        try:
            rows = json.loads(self.store.read_artifact(dataset_id, "records.json"))
            if not isinstance(rows, list):
                raise ValueError
            return rows
        except (ValueError, TypeError) as error:
            raise _error(
                "The stored dataset records are invalid.", "STORAGE_CORRUPT", 409
            ) from error
