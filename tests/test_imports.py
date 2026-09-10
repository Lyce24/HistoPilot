"""Real table imports preserve identity and reject incomplete or changed evidence."""

import base64
import json
import os
from datetime import date

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from histopilot.application import imports
from histopilot.application.imports import ImportService
from histopilot.schemas.imports import ImportSpec, InspectRequest, TableSource
from histopilot.storage.filesystem import FilesystemError, LocalFilesystem
from histopilot.storage.project_lock import StorageError
from histopilot.storage.scientific import ScientificStore


@pytest.fixture
def service(tmp_path):
    project = tmp_path / "experiment"
    project.mkdir()
    sources = tmp_path / "sources"
    sources.mkdir()
    store = ScientificStore(project, "project-import-tests")
    store.initialize()
    return ImportService(store, LocalFilesystem((sources,)))


def table(service, text, name="metadata.csv"):
    path = service.filesystem.roots[0] / name
    path.write_text(text, encoding="utf-8")
    return {"path": str(path)}


def draft(service, source=None, **options):
    spec = ImportSpec.model_validate(
        {
            "source": source or table(service, "Slide_ID,Label\n001,NA\n"),
            "slideIdColumn": "Slide_ID",
            "includeMissingSlides": True,
            **options,
        }
    )
    return service.store.create_draft(
        "import", "Bladder labels", {"type": "dataset-import", "spec": spec.model_dump()}
    )


def preview(service, saved):
    return service.preview(saved["id"], saved["revision"])


def freeze(service, saved, reviewed=None, operation="freeze-fixture"):
    reviewed = reviewed or preview(service, saved)
    return service.freeze(saved["id"], saved["revision"], reviewed["previewHash"], operation)


def codes(reviewed):
    return {item["code"] for item in reviewed["findings"]}


def field(key, owner="slide", dtype="text", **extra):
    return {"key": key, "sourceColumn": key, "owner": owner, "type": dtype, **extra}


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_inspection_examples_cover_complete_source_and_preserve_raw_values(service, extension):
    headers = ["Slide_ID", "WHO 2022", "Token", "Empty"]
    identifiers = ["0001", "A.b.uuid", "NA"] + [f"{index:04}" for index in range(4, 27)]
    labels = ["Low"] * 20 + ["high"] * 3 + ["NA", "", "low"]
    tokens = ["NA"] * 24 + ["", " "]
    rows = [
        [identifier, label, token, ""]
        for identifier, label, token in zip(identifiers, labels, tokens, strict=True)
    ]
    if extension == "csv":
        source = table(service, "\n".join(",".join(row) for row in [headers, *rows]) + "\n")
    else:
        path = service.filesystem.roots[0] / "examples.xlsx"
        book = Workbook()
        book.active.append(headers)
        for row in rows:
            book.active.append(row)
        book.save(path)
        source = {"path": str(path)}

    inspected = service.inspect(TableSource(**source))

    assert inspected["rowCount"] == 26
    assert len(inspected["rows"]) == 20
    assert {row["WHO 2022"] for row in inspected["rows"]} == {"Low"}
    assert inspected["columnSummaries"] == {
        "Slide_ID": {
            "examples": ["0001", "A.b.uuid", "NA", "0004"],
            "distinctCount": 26,
            "missingCount": 0,
        },
        "WHO 2022": {
            "examples": ["Low", "high", "NA", "low"],
            "distinctCount": 4,
            "missingCount": 1,
        },
        "Token": {"examples": ["NA", " "], "distinctCount": 2, "missingCount": 1},
        "Empty": {"examples": [], "distinctCount": 0, "missingCount": 26},
    }
    assert service.inspect(TableSource(**source)) == inspected


def test_raw_csv_identity_missing_tokens_and_snapshot_survive_source_loss(service):
    source = table(service, "De ID,Label,Age\n001,NA,050\nA.b.uuid,.,\nNA,,41\n")
    saved = draft(service, source, slideIdColumn="De ID")
    reviewed = preview(service, saved)
    assert reviewed["canFreeze"]
    assert reviewed["summary"]["mappedPatientCount"] == 0
    assert reviewed["summary"]["unlinkedSlideCount"] == 3
    assert [row["slideId"] for row in reviewed["records"]] == ["001", "A.b.uuid", "NA"]
    assert reviewed["records"][0]["attributes"] == {"Label": "NA", "Age": "050"}
    assert reviewed["records"][1]["attributes"] == {"Label": ".", "Age": None}
    source_path = service.filesystem.roots[0] / "metadata.csv"
    original = source_path.read_bytes()
    published = freeze(service, saved, reviewed)
    assert source_path.read_bytes() == original
    assert service.store.read_artifact(published["id"], "sources/main.csv") == original
    source_path.unlink()
    reopened = ImportService(
        ScientificStore(service.store.folder, service.store.project_id), service.filesystem
    )
    assert reopened.records(published["id"]) == reviewed["records"]
    assert reopened.store.get_draft(saved["id"])["status"] == "frozen"


def test_preview_freeze_retry_is_identical_after_sources_change(service):
    saved = draft(service)
    reviewed = preview(service, saved)
    published = freeze(service, saved, reviewed)
    table(service, "Slide_ID,Label\nchanged,changed\n")
    assert freeze(service, saved, reviewed) == published
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed, operation="different-operation")
    assert error.value.status_code == 409
    with pytest.raises(StorageError, match="different inputs"):
        service.freeze(saved["id"], saved["revision"], "0" * 64, "freeze-fixture")


def test_source_or_slide_inventory_change_requires_new_preview(service):
    slides = service.filesystem.roots[0] / "slides"
    slides.mkdir()
    (slides / "001.svs").write_bytes(b"slide fixture")
    saved = draft(service, slideRoot=str(slides))
    reviewed = preview(service, saved)
    (slides / "001.svs").write_bytes(b"changed slide fixture")
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed)
    assert error.value.code == "PREVIEW_STALE"
    assert service.store.list_datasets() == []
    reviewed = preview(service, saved)
    table(service, "Slide_ID,Label\n001,changed\n")
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed)
    assert error.value.code == "PREVIEW_STALE"


def test_saved_mapping_revision_is_authoritative(service):
    saved = draft(service)
    reviewed = preview(service, saved)
    payload = saved["payload"]
    payload["spec"]["missingValues"] = ["", "NA"]
    current = service.store.update_draft(
        saved["id"], expected_revision=1, name=saved["name"], payload=payload
    )
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed)
    assert error.value.code == "REVISION_CONFLICT"
    assert preview(service, current)["records"][0]["attributes"]["Label"] is None


def test_complete_recursive_scan_is_independent_of_picker_preview_limit(service):
    slides = service.filesystem.roots[0] / "slides"
    nested = slides / "batch"
    nested.mkdir(parents=True)
    ids = [f"case.{index:04d}" for index in range(211)]
    for slide_id in ids:
        (nested / f"{slide_id}.svs").write_bytes(b"slide")
    (slides / "unmatched.svs").write_bytes(b"slide")
    source = table(service, "Slide_ID\n" + "\n".join(ids + ["missing"]) + "\n")
    saved = draft(service, source, slideRoot=str(slides), includeMissingSlides=False)
    reviewed = preview(service, saved)
    assert reviewed["canFreeze"]
    assert reviewed["summary"] == {
        "sourceRowCount": 212,
        "slideCount": 211,
        "mappedPatientCount": 0,
        "verifiedPatientCount": 0,
        "fallbackSlideCount": 0,
        "unlinkedSlideCount": 211,
        "matchedSlideCount": 211,
        "missingSlideCount": 1,
        "unmatchedFileCount": 1,
        "excludedRowCount": 1,
        "scannedFileCount": 212,
    }
    assert len(reviewed["records"]) == 200
    assert reviewed["recordsTruncated"]
    published = freeze(service, saved, reviewed)
    assert len(service.records(published["id"])) == 211
    assert len(json.loads(service.store.read_artifact(published["id"], "inventory.json"))) == 212
    exclusions = json.loads(service.store.read_artifact(published["id"], "exclusions.json"))
    assert exclusions[0]["slideId"] == "missing"


def test_metadata_only_requires_explicit_inclusion_rule(service):
    saved = draft(service, includeMissingSlides=False)
    reviewed = preview(service, saved)
    assert not reviewed["canFreeze"]
    assert "SLIDE_SOURCE_REQUIRED" in codes(reviewed)
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed)
    assert error.value.code == "IMPORT_BLOCKED"


@pytest.mark.parametrize("rows,expected", [("a\na\n", "SLIDE_ID_DUPLICATE"), ("a\nb,c\n", "width")])
def test_duplicate_or_malformed_metadata_is_not_frozen(service, rows, expected):
    saved = draft(service, table(service, "Slide_ID\n" + rows))
    if expected == "width":
        with pytest.raises(StorageError):
            preview(service, saved)
    else:
        reviewed = preview(service, saved)
        assert expected in codes(reviewed)
        assert not reviewed["canFreeze"]


@pytest.mark.parametrize("alias", ["stem", "hardlink", "symlink"])
def test_duplicate_slide_stems_and_physical_aliases_block(service, alias):
    slides = service.filesystem.roots[0] / "slides"
    slides.mkdir()
    original = slides / "001.svs"
    original.write_bytes(b"slide")
    if alias == "stem":
        (slides / "001.tif").write_bytes(b"different")
        expected = "SLIDE_MATCH_AMBIGUOUS"
    elif alias == "hardlink":
        os.link(original, slides / "002.svs")
        expected = "DUPLICATE_SLIDE_ALIAS"
    else:
        (slides / "002.svs").symlink_to(original)
        expected = "DUPLICATE_SLIDE_ALIAS"
    reviewed = preview(service, draft(service, slideRoot=str(slides)))
    assert not reviewed["canFreeze"]
    assert expected in codes(reviewed)


def test_sources_and_discovered_links_cannot_escape_allowed_data_roots(service, tmp_path):
    outside = tmp_path / "outside.csv"
    outside.write_text("Slide_ID\nsecret\n")
    with pytest.raises(FilesystemError) as error:
        service.inspect(TableSource(path=str(outside)))
    assert error.value.status_code == 403
    slides = service.filesystem.roots[0] / "slides"
    slides.mkdir()
    (slides / "001.svs").symlink_to(outside)
    reviewed = preview(service, draft(service, slideRoot=str(slides)))
    assert "SOURCE_PATH_ESCAPE" in codes(reviewed)
    assert not reviewed["canFreeze"]


@pytest.mark.parametrize("bound", ["MAX_FILES", "MAX_ENTRIES"])
def test_scan_overflow_refuses_partial_inventory(service, monkeypatch, bound):
    slides = service.filesystem.roots[0] / "slides"
    slides.mkdir()
    for index in range(3):
        (slides / f"{index}.svs").write_bytes(b"slide")
    monkeypatch.setattr(imports, bound, 2)
    saved = draft(service, slideRoot=str(slides))
    with pytest.raises(StorageError) as error:
        preview(service, saved)
    assert error.value.code == "SCAN_LIMIT"
    assert service.store.list_datasets() == []


def test_table_overflow_refuses_partial_rows(service, monkeypatch):
    monkeypatch.setattr(imports, "MAX_ROWS", 2)
    with pytest.raises(StorageError) as error:
        service.inspect(TableSource(**table(service, "Slide_ID\na\nb\nc\n")))
    assert error.value.code == "TABLE_ROW_LIMIT"


def test_join_amplification_is_bounded_before_whole_result_serialization(service, monkeypatch):
    source = table(service, "Slide_ID,Patient_ID\ns1,p1\ns2,p1\ns3,p1\n")
    patients = table(service, "Patient_ID,Notes\np1," + "x" * 300 + "\n", "patients.csv")
    saved = draft(
        service,
        source,
        patientIdColumn="Patient_ID",
        patientSource=patients,
        patientSourceKind="patients",
        patientAttributes=[field("Notes", "patient")],
    )
    monkeypatch.setattr(imports, "MAX_ARTIFACT_BYTES", 700)
    original_json = imports._json

    def bounded_json(value):
        assert not (isinstance(value, dict) and "records" in value)
        return original_json(value)

    monkeypatch.setattr(imports, "_json", bounded_json)
    with pytest.raises(StorageError) as error:
        preview(service, saved)
    assert error.value.code == "IMPORT_TOO_LARGE"
    assert error.value.status_code == 413
    assert service.store.list_datasets() == []


def test_xlsx_sheets_exact_string_identifiers_and_formulas(service):
    path = service.filesystem.roots[0] / "bladder.xlsx"
    book = Workbook()
    book.active.title = "Instructions"
    book.active.append(["Read me"])
    book.active.append(["Select clinical data"])
    sheet = book.create_sheet("Clinical data")
    sheet.append(["De ID", "Label", "Age", "Formula"])
    sheet.append(["0001", "NA", 51, "=1+1"])
    sheet.append(["A.b", ".", 60, "literal"])
    book.save(path)
    source = {"path": str(path), "sheet": "Clinical data"}
    inspected = service.inspect(InspectRequest(source=TableSource(**source)))
    assert inspected["sheets"] == ["Instructions", "Clinical data"]
    assert inspected["rows"][0] == {"De ID": "0001", "Label": "NA", "Age": "51", "Formula": "=1+1"}
    saved = draft(service, source, slideIdColumn="De ID")
    assert "MAPPED_FORMULA" in codes(preview(service, saved))
    saved = draft(
        service,
        source,
        slideIdColumn="De ID",
        attributes=[field("Label"), field("Age", dtype="integer")],
    )
    reviewed = preview(service, saved)
    assert reviewed["canFreeze"]
    assert service.records(freeze(service, saved, reviewed)["id"])[0]["slideId"] == "0001"


def test_xlsx_numeric_identifiers_have_explicit_warning(service):
    path = service.filesystem.roots[0] / "numeric.xlsx"
    book = Workbook()
    book.active.append(["Slide_ID"])
    book.active.append([1])
    book.active["A2"].number_format = "0000"
    book.save(path)
    reviewed = preview(service, draft(service, {"path": str(path)}))
    assert reviewed["records"][0]["slideId"] == "1"
    assert "NUMERIC_IDENTIFIER" in codes(reviewed)


def test_xlsx_dates_preserve_the_stored_temporal_value(service):
    path = service.filesystem.roots[0] / "dates.xlsx"
    book = Workbook()
    book.active.append(["Slide_ID", "Collected"])
    book.active.append(["001", date(2024, 2, 29)])
    book.save(path)
    reviewed = preview(
        service,
        draft(service, {"path": str(path)}, attributes=[field("Collected", dtype="date")]),
    )
    assert reviewed["canFreeze"]
    assert reviewed["records"][0]["attributes"]["Collected"] == "2024-02-29T00:00:00"


@pytest.mark.parametrize(
    "content,code",
    [
        ("Slide_ID,Slide_ID\na,b\n", "TABLE_HEADERS_DUPLICATE"),
        ("Slide_ID,\na,b\n", "TABLE_HEADERS_INVALID"),
        ("Slide_ID\n", "TABLE_EMPTY"),
    ],
)
def test_invalid_table_headers_and_empty_tables_rejected(service, content, code):
    with pytest.raises(StorageError) as error:
        service.inspect(TableSource(**table(service, content)))
    assert error.value.code == code


def test_upload_snapshots_are_bounded_and_never_use_user_filename_as_path(service):
    content = b"Slide_ID,Label\n001,NA\n"
    source = {"contentBase64": base64.b64encode(content).decode(), "filename": "../../data.csv"}
    saved = draft(service, source)
    reviewed = preview(service, saved)
    published = freeze(service, saved, reviewed)
    assert service.store.read_artifact(published["id"], "sources/main.csv") == content
    assert "contentBase64" not in published["manifest"]["provenance"]["mapping"]["source"]
    with pytest.raises(StorageError) as error:
        service.inspect(TableSource(contentBase64="!bad!", filename="bad.csv"))
    assert error.value.code == "UPLOAD_INVALID"
    with pytest.raises(StorageError) as error:
        service.inspect(
            TableSource(
                contentBase64=base64.b64encode(b"x" * (imports.MAX_UPLOAD_BYTES + 1)).decode(),
                filename="large.csv",
            )
        )
    assert error.value.code == "UPLOAD_TOO_LARGE"
    with pytest.raises(ValidationError):
        TableSource(path="/tmp/example.csv", **source)


@pytest.mark.parametrize(
    "values,code",
    [("60\n61", "PATIENT_ATTRIBUTE_CONFLICT"), ("60\n", "PATIENT_ATTRIBUTE_INCOMPLETE")],
)
def test_patient_owned_conflicting_or_partly_missing_values_block(service, values, code):
    first, second = values.split("\n")
    source = table(service, f"Slide_ID,Patient_ID,Age,Label\ns1,p1,{first},A\ns2,p1,{second},B\n")
    saved = draft(
        service,
        source,
        patientIdColumn="Patient_ID",
        attributes=[field("Age", "patient", "integer"), field("Label")],
    )
    reviewed = preview(service, saved)
    assert not reviewed["canFreeze"]
    assert code in codes(reviewed)
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed)
    assert error.value.code == "IMPORT_BLOCKED"


def test_patient_ownership_and_slide_label_heterogeneity_are_independent(service):
    source = table(service, "Slide_ID,Patient_ID,Age,Label\ns1,p1,,A\ns2,p1,,B\n")
    saved = draft(
        service,
        source,
        patientIdColumn="Patient_ID",
        attributes=[field("Age", "patient", "integer"), field("Label", dtype="categorical")],
    )
    reviewed = preview(service, saved)
    assert reviewed["canFreeze"]
    assert reviewed["summary"]["mappedPatientCount"] == 1
    assert [row["attributes"]["Label"] for row in reviewed["records"]] == ["A", "B"]
    assert all(row["attributes"]["Age"] is None for row in reviewed["records"])


@pytest.mark.parametrize("fallback", [None, "unresolved", "slide_id"])
def test_patient_fallback_is_explicit_and_only_fills_missing_ids(service, fallback):
    source = table(service, "Slide_ID,Patient_ID,Label\ns1,p1,A\ns2,p1,A\ns3,,B\ns4,unknown,B\n")
    options = {"patientIdFallback": fallback} if fallback is not None else {}
    saved = draft(
        service,
        source,
        patientIdColumn="Patient_ID",
        missingValues=["", "unknown"],
        **options,
    )
    reviewed = preview(service, saved)
    assert reviewed["canFreeze"]
    assert reviewed["summary"]["verifiedPatientCount"] == 1
    assert [row["patientId"] for row in reviewed["records"][:2]] == ["p1", "p1"]
    assert [row["patientIdSource"] for row in reviewed["records"][:2]] == ["source", "source"]
    if fallback == "slide_id":
        assert [row["patientId"] for row in reviewed["records"][2:]] == ["s3", "s4"]
        assert all(row["patientIdSource"] == "slide_fallback" for row in reviewed["records"][2:])
        assert reviewed["summary"]["mappedPatientCount"] == 3
        assert reviewed["summary"]["fallbackSlideCount"] == 2
        assert reviewed["summary"]["unlinkedSlideCount"] == 0
        warning = next(
            item for item in reviewed["findings"] if item["code"] == "PATIENT_ID_SLIDE_FALLBACK"
        )
        assert warning["severity"] == "warning"
        assert warning["count"] == 2
        assert "not verified" in warning["message"]
        assert "PATIENT_ID_UNRESOLVED" not in codes(reviewed)
    else:
        assert all(row["patientId"] is None for row in reviewed["records"][2:])
        assert all(row["patientIdSource"] == "unresolved" for row in reviewed["records"][2:])
        assert reviewed["summary"]["mappedPatientCount"] == 1
        assert reviewed["summary"]["fallbackSlideCount"] == 0
        assert reviewed["summary"]["unlinkedSlideCount"] == 2
        assert "PATIENT_ID_SLIDE_FALLBACK" not in codes(reviewed)


def test_crosswalk_resolves_patient_before_slide_fallback(service):
    source = table(service, "Slide_ID,Patient_ID\ns1,p1\ns2,\ns3,\n")
    crosswalk = table(service, "Slide_ID,Patient_ID\ns1,p1\ns2,p1\n", "crosswalk.csv")
    reviewed = preview(
        service,
        draft(
            service,
            source,
            patientIdColumn="Patient_ID",
            patientSource=crosswalk,
            patientIdFallback="slide_id",
        ),
    )
    assert reviewed["canFreeze"]
    assert [row["patientId"] for row in reviewed["records"]] == ["p1", "p1", "s3"]
    assert [row["patientIdSource"] for row in reviewed["records"]] == [
        "crosswalk",
        "crosswalk",
        "slide_fallback",
    ]
    assert reviewed["summary"]["verifiedPatientCount"] == 1
    assert reviewed["summary"]["fallbackSlideCount"] == 1


def test_fallback_slide_id_does_not_join_patient_attribute_table(service):
    source = table(service, "Slide_ID,Patient_ID\ns1,p1\ns2,\n")
    patients = table(service, "Patient_ID,Age\np1,60\ns2,99\n", "patients.csv")
    reviewed = preview(
        service,
        draft(
            service,
            source,
            patientIdColumn="Patient_ID",
            patientSource=patients,
            patientSourceKind="patients",
            patientAttributes=[field("Age", "patient", "integer")],
            patientIdFallback="slide_id",
        ),
    )
    assert reviewed["records"][1]["patientId"] == "s2"
    assert reviewed["records"][1]["patientIdSource"] == "slide_fallback"
    assert reviewed["records"][1]["attributes"]["Age"] is None
    assert "PATIENT_SOURCE_UNMATCHED" in codes(reviewed)


@pytest.mark.parametrize("use_crosswalk", [False, True])
def test_fallback_cannot_merge_a_slide_with_an_existing_patient_id(service, use_crosswalk):
    source = table(service, "Slide_ID,Patient_ID\ns1,\ns2,s1\n")
    options = {"patientIdColumn": "Patient_ID"}
    if use_crosswalk:
        options = {"patientSource": table(service, "Slide_ID,Patient_ID\ns2,s1\n", "crosswalk.csv")}
    saved = draft(service, source, patientIdFallback="slide_id", **options)
    reviewed = preview(service, saved)
    assert not reviewed["canFreeze"]
    collision = next(
        item for item in reviewed["findings"] if item["code"] == "PATIENT_ID_FALLBACK_COLLISION"
    )
    assert collision["severity"] == "error"
    assert collision["examples"] == ["s1"]
    with pytest.raises(StorageError) as error:
        freeze(service, saved, reviewed)
    assert error.value.code == "IMPORT_BLOCKED"


def test_fallback_acknowledgement_and_identity_sources_survive_freeze_and_reopen(service):
    source = table(service, "Slide_ID,Patient_ID\ns1,p1\ns2,\n")
    unresolved = draft(service, source, patientIdColumn="Patient_ID")
    saved = draft(service, source, patientIdColumn="Patient_ID", patientIdFallback="slide_id")
    reviewed = preview(service, saved)
    assert preview(service, unresolved)["previewHash"] != reviewed["previewHash"]
    published = freeze(service, saved, reviewed)
    source_path = service.filesystem.roots[0] / "metadata.csv"
    source_path.unlink()
    reopened = ImportService(
        ScientificStore(service.store.folder, service.store.project_id), service.filesystem
    )
    assert reopened.records(published["id"]) == reviewed["records"]
    manifest = reopened.store.get_dataset(published["id"])["manifest"]
    assert manifest["summary"]["verifiedPatientCount"] == 1
    assert manifest["summary"]["fallbackSlideCount"] == 1
    assert manifest["provenance"]["mapping"]["patientIdFallback"] == "slide_id"
    saved_spec = reopened.store.get_draft(saved["id"])["payload"]["spec"]
    assert saved_spec["patientIdFallback"] == "slide_id"
    provenance = json.loads(reopened.store.read_artifact(published["id"], "provenance.json"))
    assert provenance["mapping"]["patientIdFallback"] == "slide_id"


def test_late_crosswalk_creates_linked_child_without_changing_frozen_parent(service):
    source = table(service, "Slide_ID,Label\ns1,A\ns2,B\n")
    original = draft(service, source)
    parent = freeze(service, original)
    crosswalk = table(service, "Slide_ID,Patient_ID,Age\ns1,p001,60\ns2,p001,60\n", "crosswalk.csv")
    child_draft = draft(
        service,
        source,
        parentId=parent["id"],
        patientSource=crosswalk,
        patientAttributes=[field("Age", "patient", "integer")],
    )
    reviewed = preview(service, child_draft)
    assert reviewed["canFreeze"]
    assert reviewed["summary"]["mappedPatientCount"] == 1
    child = freeze(service, child_draft, reviewed, operation="late-crosswalk")
    assert child["parentId"] == parent["id"]
    assert child["id"] != parent["id"]
    assert all(row["patientId"] is None for row in service.records(parent["id"]))
    assert {row["patientId"] for row in service.records(child["id"])} == {"p001"}


def test_patient_keyed_attribute_table_joins_only_explicit_patient_ids(service):
    source = table(service, "Slide_ID,Patient_ID\ns1,p1\ns2,p1\n")
    patients = table(service, "Patient_ID,Age\np1,60\np2,70\n", "patients.csv")
    options = {
        "patientSource": patients,
        "patientSourceKind": "patients",
        "patientAttributes": [field("Age", "patient", "integer")],
    }
    saved = draft(service, source, patientIdColumn="Patient_ID", **options)
    reviewed = preview(service, saved)
    assert reviewed["canFreeze"]
    assert [row["attributes"]["Age"] for row in reviewed["records"]] == ["60", "60"]
    assert "PATIENT_SOURCE_UNMATCHED" in codes(reviewed)
    missing_join = draft(service, source, attributes=[], **options)
    assert "PATIENT_JOIN_REQUIRED" in codes(preview(service, missing_join))


@pytest.mark.parametrize("crosswalk", [False, True])
def test_slide_identity_cannot_implicitly_become_patient_identity(service, crosswalk):
    source = table(service, "Slide_ID\n001\n")
    options = (
        {
            "patientSource": source,
            "patientSourceSlideIdColumn": "Slide_ID",
            "patientSourcePatientIdColumn": "Slide_ID",
        }
        if crosswalk
        else {"patientIdColumn": "Slide_ID"}
    )
    reviewed = preview(service, draft(service, source, **options))
    assert "PATIENT_MAPPING_REQUIRES_CROSSWALK" in codes(reviewed)
    assert not reviewed["canFreeze"]


@pytest.mark.parametrize(
    "patient_rows,code",
    [
        ("s1,p1\ns1,p1\n", "PATIENT_SOURCE_KEY_DUPLICATE"),
        ("s1,\n", "PATIENT_SOURCE_KEY_MISSING"),
        ("s1,p2\n", "PATIENT_LINK_CONFLICT"),
    ],
)
def test_ambiguous_missing_and_disagreeing_crosswalks_block(service, patient_rows, code):
    source = table(service, "Slide_ID,Patient_ID\ns1,p1\n")
    crosswalk = table(service, "Slide_ID,Patient_ID\n" + patient_rows, "crosswalk.csv")
    reviewed = preview(
        service, draft(service, source, patientIdColumn="Patient_ID", patientSource=crosswalk)
    )
    assert code in codes(reviewed)
    assert not reviewed["canFreeze"]


@pytest.mark.parametrize(
    "dtype,value",
    [
        ("integer", "1.1"),
        ("decimal", "NaN"),
        ("decimal", "1e1000"),
        ("boolean", "unknown"),
        ("date", "yesterday"),
        ("categorical", "C"),
    ],
)
def test_explicit_attribute_types_block_invalid_values(service, dtype, value):
    source = table(service, f"Slide_ID,Value\ns1,{value}\n")
    selected = field(
        "Value", dtype=dtype, categories=["A", "B"] if dtype == "categorical" else None
    )
    reviewed = preview(service, draft(service, source, attributes=[selected]))
    assert "ATTRIBUTE_TYPE_INVALID" in codes(reviewed)
    assert not reviewed["canFreeze"]


def test_empty_attribute_selection_and_per_field_missing_policy_are_explicit(service):
    source = table(service, "Slide_ID,Label\n001,NA\n")
    assert preview(service, draft(service, source, attributes=[]))["dictionary"] == []
    saved = draft(service, source, attributes=[field("Label", missingValues=["NA"])])
    assert preview(service, saved)["records"][0]["attributes"]["Label"] is None


def test_long_source_header_requires_explicit_short_attribute_key(service):
    column = "x" * 129
    source = table(service, f"Slide_ID,{column}\n001,A\n")
    with pytest.raises(StorageError) as error:
        preview(service, draft(service, source))
    assert error.value.code == "ATTRIBUTE_KEY_REQUIRED"
    saved = draft(service, source, attributes=[{"key": "Label", "sourceColumn": column}])
    assert preview(service, saved)["records"][0]["attributes"] == {"Label": "A"}


@pytest.mark.parametrize(
    "kind,payload",
    [
        ("experiment", {"type": "dataset-import", "spec": {}}),
        ("import", {"status": "ready", "records": []}),
        ("import", {"type": "dataset-import", "spec": {}, "canFreeze": True}),
    ],
)
def test_user_draft_payload_cannot_assert_import_readiness(service, kind, payload):
    saved = service.store.create_draft(kind, "Untrusted intent", payload)
    with pytest.raises(StorageError) as error:
        preview(service, saved)
    assert error.value.code == "IMPORT_DRAFT_REQUIRED"
    assert service.store.list_datasets() == []
